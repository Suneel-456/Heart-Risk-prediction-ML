"""Build data/heart_disease_data.xlsx (data, dictionary, model results).

Run:  python -m src.export_excel
"""
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config
from src.data_preprocessing import load_raw

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style(ws, df):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)


def main():
    data = load_raw()
    dictionary = pd.DataFrame(
        [(k, v[0], v[1]) for k, v in config.FEATURE_DESCRIPTIONS.items()],
        columns=["column", "description", "type"],
    )
    summary = data.describe().T.round(2).reset_index().rename(columns={"index": "column"})
    sheets = {"Data": data, "Data_Dictionary": dictionary, "Summary_Stats": summary}

    comp = config.REPORTS_DIR / "model_comparison.csv"
    if comp.exists():
        sheets["Model_Comparison"] = pd.read_csv(comp).drop(columns=["best_params"]).round(3)
    pred = config.REPORTS_DIR / "test_predictions.csv"
    if pred.exists():
        sheets["Test_Predictions"] = pd.read_csv(pred).rename(columns={"Unnamed: 0": "row_id"})
    imp = config.REPORTS_DIR / "feature_importance.csv"
    if imp.exists():
        sheets["Feature_Importance"] = pd.read_csv(imp).rename(
            columns={"Unnamed: 0": "feature"}).round(4)

    config.EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(config.EXCEL_PATH, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            _style(xw.sheets[name], df)
    print(f"Wrote {config.EXCEL_PATH} with sheets: {', '.join(sheets)}")


if __name__ == "__main__":
    main()
